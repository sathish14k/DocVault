"""Text extraction for search indexing. Never raises: returns '' on failure."""
import os
import subprocess
import tempfile
from pathlib import Path

OCR_LANGS = os.environ.get("OCR_LANGS", "eng")
IMG_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp", ".gif"}
TEXT_EXT = {".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".log", ".rtf"}
LIBRE_EXT = {".doc", ".odt", ".ppt", ".pptx", ".odp", ".xls", ".ods"}
MAX_CHARS = 2_000_000


def ocr_image(path: Path) -> str:
    r = subprocess.run(["tesseract", str(path), "stdout", "-l", OCR_LANGS],
                       capture_output=True, text=True, timeout=300)
    return r.stdout


def ocr_pdf(path: Path, max_pages: int = 40) -> str:
    with tempfile.TemporaryDirectory() as t:
        subprocess.run(["pdftoppm", "-r", "200", "-l", str(max_pages), "-png", str(path), f"{t}/p"],
                       check=True, timeout=900, capture_output=True)
        return "\n".join(ocr_image(p) for p in sorted(Path(t).glob("p*.png")))


def pdf_text(path: Path) -> str:
    from pypdf import PdfReader
    r = PdfReader(str(path))
    if r.is_encrypted:
        r.decrypt("")
    text = "\n".join((p.extract_text() or "") for p in r.pages)
    if len(text.strip()) < 40:  # likely a scan
        try:
            text = ocr_pdf(path)
        except Exception:
            pass
    return text


def docx_text(path: Path) -> str:
    import docx
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def xlsx_text(path: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"# {ws.title}")
        for row in ws.iter_rows(values_only=True):
            vals = [str(v) for v in row if v is not None]
            if vals:
                parts.append(" | ".join(vals))
    return "\n".join(parts)


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            text = pdf_text(path)
        elif ext == ".docx":
            text = docx_text(path)
        elif ext in (".xlsx", ".xlsm"):
            text = xlsx_text(path)
        elif ext in TEXT_EXT:
            text = path.read_text(errors="ignore")
        elif ext in IMG_EXT:
            text = ocr_image(path)
        elif ext in LIBRE_EXT:
            from .main import as_pdf  # converts via LibreOffice, cached
            text = pdf_text(as_pdf(0, path, cache_name=f"tmp-{path.stem}"))
        else:
            text = ""
    except Exception:
        text = ""
    return text[:MAX_CHARS]
